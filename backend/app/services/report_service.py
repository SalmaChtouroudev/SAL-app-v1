from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
import unicodedata

from openpyxl import Workbook

from app.core.config import Settings
from app.repositories.platform_repository import PlatformRepository
from app.schemas.platform import (
    AuditLogCreate,
    ReportAlertDigestRead,
    ReportBreakdownItem,
    ReportExportRead,
    ReportExportRequest,
    ReportExportResponse,
    ReportPreviewRead,
)


@dataclass(frozen=True)
class ReportMetric:
    label: str
    value: str


@dataclass(frozen=True)
class ReportInsight:
    title: str
    description: str


@dataclass(frozen=True)
class ReportSnapshotData:
    company_name: str
    generated_at: str
    reporting_period: str
    latest_bucket: str | None
    total_revenue: float
    total_expenses: float
    net_margin: float
    growth_percent: float
    charge_ratio: float
    active_alerts: int
    critical_alerts: int
    priority_recommendation: str
    alerts_digest: list[ReportAlertDigestRead]
    metrics: list[ReportMetric]
    top_products: list[ReportBreakdownItem]
    top_clients: list[ReportBreakdownItem]
    insights: list[ReportInsight]
    trend_rows: list[tuple[str, float, float, float]]


@dataclass(frozen=True)
class PdfColor:
    r: float
    g: float
    b: float


class PdfDocumentBuilder:
    def __init__(self) -> None:
        self.commands: list[str] = []

    def draw_rect(
        self,
        x: float,
        y: float,
        width: float,
        height: float,
        *,
        fill: PdfColor | None = None,
        stroke: PdfColor | None = None,
        line_width: float = 1.0,
    ) -> None:
        if fill is not None:
            self.commands.append(f"{fill.r} {fill.g} {fill.b} rg")
        if stroke is not None:
            self.commands.append(f"{stroke.r} {stroke.g} {stroke.b} RG")
            self.commands.append(f"{line_width} w")

        self.commands.append(f"{x:.2f} {y:.2f} {width:.2f} {height:.2f} re")
        self.commands.append("B" if fill is not None and stroke is not None else "f" if fill else "S")

    def draw_text(
        self,
        text: str,
        x: float,
        y: float,
        *,
        font: str = "F1",
        size: float = 12,
        color: PdfColor | None = None,
    ) -> None:
        normalized = escape_pdf_text(text)
        if not normalized:
            return

        tone = color or PDF_COLORS["ink"]
        self.commands.extend(
            [
                "BT",
                f"/{font} {size} Tf",
                f"{tone.r} {tone.g} {tone.b} rg",
                f"1 0 0 1 {x:.2f} {y:.2f} Tm",
                f"({normalized}) Tj",
                "ET",
            ]
        )

    def draw_text_block(
        self,
        text: str,
        x: float,
        top: float,
        width: float,
        *,
        font: str = "F1",
        size: float = 12,
        color: PdfColor | None = None,
        leading: float | None = None,
        max_lines: int = 3,
    ) -> float:
        step = leading if leading is not None else size + 4
        baseline = top - size

        for line in wrap_pdf_text(text, width, size, max_lines):
            self.draw_text(line, x, baseline, font=font, size=size, color=color)
            baseline -= step

        return baseline + step

    def draw_right_text(
        self,
        text: str,
        right: float,
        y: float,
        *,
        font: str = "F1",
        size: float = 12,
        color: PdfColor | None = None,
    ) -> None:
        self.draw_text(text, right - estimate_text_width(text, size), y, font=font, size=size, color=color)

    def build_bytes(self) -> bytes:
        content = "\n".join(self.commands)
        objects = [
            "1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj",
            "2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj",
            "3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R /F2 6 0 R >> >> >> endobj",
            f"4 0 obj << /Length {len(content)} >> stream\n{content}\nendstream endobj",
            "5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj",
            "6 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >> endobj",
        ]

        pdf = "%PDF-1.4\n"
        offsets: list[int] = []
        for object_block in objects:
            offsets.append(len(pdf))
            pdf += f"{object_block}\n"

        xref_offset = len(pdf)
        pdf += "xref\n"
        pdf += f"0 {len(offsets) + 1}\n"
        pdf += "0000000000 65535 f \n"
        for offset in offsets:
            pdf += f"{offset:010d} 00000 n \n"
        pdf += f"trailer << /Size {len(offsets) + 1} /Root 1 0 R >>\n"
        pdf += f"startxref\n{xref_offset}\n%%EOF"
        return pdf.encode("ascii")


PDF_COLORS: dict[str, PdfColor] = {
    "brand": PdfColor(0.08, 0.68, 0.84),
    "brand_dark": PdfColor(0.06, 0.45, 0.58),
    "ink": PdfColor(0.08, 0.16, 0.25),
    "muted": PdfColor(0.39, 0.46, 0.53),
    "panel": PdfColor(0.95, 0.97, 0.99),
    "border": PdfColor(0.83, 0.88, 0.92),
    "white": PdfColor(1.0, 1.0, 1.0),
    "danger": PdfColor(0.82, 0.21, 0.25),
    "warning": PdfColor(0.93, 0.55, 0.12),
    "success": PdfColor(0.09, 0.56, 0.41),
}


def sanitize_pdf_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_bytes = normalized.encode("ascii", "ignore")
    return " ".join(ascii_bytes.decode("ascii").split())


def escape_pdf_text(value: str) -> str:
    return sanitize_pdf_text(value).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def estimate_text_width(value: str, size: float) -> float:
    return len(sanitize_pdf_text(value)) * size * 0.5


def wrap_pdf_text(value: str, width: float, size: float, max_lines: int = 3) -> list[str]:
    words = sanitize_pdf_text(value).split()
    if not words:
        return [""]

    lines: list[str] = []
    current = ""

    for word in words:
        candidate = f"{current} {word}".strip()
        if estimate_text_width(candidate, size) <= width:
            current = candidate
            continue

        if current:
            lines.append(current)
            current = word
        else:
            slice_length = max(1, int(width / (size * 0.5)) - 1)
            lines.append(f"{word[:slice_length]}-")
            current = word[slice_length:]

        if len(lines) >= max_lines:
            break

    if current and len(lines) < max_lines:
        lines.append(current)

    if len(lines) > max_lines:
        lines = lines[:max_lines]

    if len(lines) == max_lines and words:
        last = lines[-1]
        lines[-1] = f"{last[:-3]}..." if len(last) > 3 else f"{last}..."

    return lines


class ReportService:
    def __init__(self, repository: PlatformRepository, settings: Settings) -> None:
        self.repository = repository
        self.settings = settings
        self.storage_dir = Path(settings.report_storage_dir)

    def export(self, payload: ReportExportRequest) -> ReportExportResponse:
        snapshot = self.build_snapshot(payload.companyId)
        extension = "pdf" if payload.format == "pdf" else "xlsx"
        company_slug = snapshot.company_name.replace(" ", "-").replace("/", "-").lower()
        file_name = f"{company_slug}-executive-report.{extension}"
        export = self.repository.record_report_export(
            company_id=payload.companyId,
            requested_by=payload.requestedBy,
            format_name=payload.format,
            file_name=file_name,
            status="COMPLETED",
        )
        output_path = self.build_output_path(export)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        if payload.format == "pdf":
            output_path.write_bytes(self.build_pdf(snapshot))
        else:
            output_path.write_bytes(self.build_excel(snapshot))

        self.repository.create_audit_log(
            AuditLogCreate(
                enterpriseId=payload.companyId,
                userId=payload.requestedBy,
                actionType="REPORT_EXPORTED",
                status="SUCCESS",
                description=f"{payload.format.upper()} report generated on the server: {export.fileName}.",
            )
        )

        return ReportExportResponse(
            id=export.id,
            status="COMPLETED",
            fileName=export.fileName,
            message="Report generated successfully and is ready to download.",
            generatedAt=snapshot.generated_at,
            downloadPath=self.build_download_path(export.id),
        )

    def list_export_history(self, company_id: str) -> list[ReportExportRead]:
        return [
            export.model_copy(update={"downloadPath": self.build_download_path(export.id)})
            for export in self.repository.list_report_exports(company_id)
        ]

    def get_export(self, export_id: str) -> ReportExportRead | None:
        export = self.repository.get_report_export(export_id)
        if export is None:
            return None
        return export.model_copy(update={"downloadPath": self.build_download_path(export.id)})

    def resolve_output_path(self, export: ReportExportRead) -> Path:
        return self.build_output_path(export)

    def build_download_path(self, export_id: str) -> str:
        return f"/api/reports/{export_id}/download"

    def preview(self, company_id: str) -> ReportPreviewRead:
        snapshot = self.build_snapshot(company_id)
        return ReportPreviewRead(
            companyName=snapshot.company_name,
            generatedAt=snapshot.generated_at,
            totalRevenue=snapshot.total_revenue,
            totalExpenses=snapshot.total_expenses,
            netMargin=snapshot.net_margin,
            growthPercent=snapshot.growth_percent,
            chargeRatio=snapshot.charge_ratio,
            activeAlerts=snapshot.active_alerts,
            criticalAlerts=snapshot.critical_alerts,
            priorityRecommendation=snapshot.priority_recommendation,
            alertsDigest=snapshot.alerts_digest,
            topProducts=snapshot.top_products,
            topClients=snapshot.top_clients,
        )

    def build_snapshot(self, company_id: str) -> ReportSnapshotData:
        company = self.repository.get_company(company_id)
        if company is None:
            raise ValueError(f"Unknown company '{company_id}'.")

        trend = self.repository.get_dashboard_trend(company_id)
        total_revenue = round(sum(point.totalRevenue for point in trend), 2)
        total_expenses = round(sum(point.totalExpenses for point in trend), 2)
        net_margin = round(total_revenue - total_expenses, 2)
        latest_bucket = trend[-1].bucket if trend else None
        previous_revenue = trend[-2].totalRevenue if len(trend) > 1 else 0
        latest_revenue = trend[-1].totalRevenue if trend else 0
        growth_percent = (
            round(((latest_revenue - previous_revenue) / previous_revenue) * 100, 2)
            if previous_revenue
            else 0.0
        )
        charge_ratio = round((total_expenses / total_revenue) * 100, 2) if total_revenue else 0.0
        active_alerts = [alert for alert in self.repository.list_alerts(company_id) if alert.isActive]
        critical_alerts = sum(1 for alert in active_alerts if alert.severity == "critical")
        alerts_digest = [
            ReportAlertDigestRead(
                alertId=alert.id,
                ruleId=alert.ruleId,
                ruleName=alert.ruleName,
                severity=alert.severity,
                title=alert.title,
                recommendation=alert.recommendation,
                status="ACTIVE" if alert.isActive else "RESOLVED",
                createdAt=alert.createdAt,
            )
            for alert in active_alerts[:4]
        ]
        priority_recommendation = (
            active_alerts[0].recommendation
            if active_alerts and active_alerts[0].recommendation.strip()
            else (
                "Review the latest KPI trend and focus on the weakest operating signals first."
            )
        )

        insights = [
            ReportInsight(
                title="Margin watch",
                description=(
                    "Margins remain under pressure. Prioritize cost reviews and pricing discipline."
                    if net_margin <= 25000
                    else "Margins remain healthy. Continue scaling the best-performing offers."
                ),
            ),
            ReportInsight(
                title="Growth pulse",
                description=(
                    "Revenue momentum is softening. Strengthen upsell and retention actions."
                    if growth_percent < 5
                    else "Growth remains on track. Double down on the strongest client segments."
                ),
            ),
            ReportInsight(
                title="Cost discipline",
                description=(
                    "Charges are above the safe operating band. Review recurring expenses immediately."
                    if charge_ratio > 70
                    else "Charges remain inside an acceptable operating corridor."
                ),
            ),
        ]

        return ReportSnapshotData(
            company_name=company.name,
            generated_at=datetime.now(tz=timezone.utc).isoformat(),
            reporting_period="Rolling 6 months",
            latest_bucket=latest_bucket,
            total_revenue=total_revenue,
            total_expenses=total_expenses,
            net_margin=net_margin,
            growth_percent=growth_percent,
            charge_ratio=charge_ratio,
            active_alerts=len(active_alerts),
            critical_alerts=critical_alerts,
            priority_recommendation=priority_recommendation,
            alerts_digest=alerts_digest,
            metrics=[
                ReportMetric(label="Revenue", value=self.format_currency(total_revenue)),
                ReportMetric(label="Charges", value=self.format_currency(total_expenses)),
                ReportMetric(label="Margin", value=self.format_currency(net_margin)),
                ReportMetric(label="Growth", value=self.format_percent(growth_percent)),
                ReportMetric(label="Charge ratio", value=self.format_percent(charge_ratio)),
            ],
            top_products=self.repository.list_top_products(company_id, limit=5),
            top_clients=self.repository.list_top_clients(company_id, limit=5),
            insights=insights,
            trend_rows=[
                (point.bucket, point.totalRevenue, point.totalExpenses, point.netMargin)
                for point in trend
            ],
        )

    def build_excel(self, snapshot: ReportSnapshotData) -> bytes:
        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        summary_sheet.append(["Company", snapshot.company_name])
        summary_sheet.append(["Generated at", snapshot.generated_at])
        summary_sheet.append(["Reporting period", snapshot.reporting_period])
        summary_sheet.append(["Latest trend bucket", snapshot.latest_bucket or "N/A"])
        summary_sheet.append([])
        summary_sheet.append(["Metric", "Value"])
        for metric in snapshot.metrics:
            summary_sheet.append([metric.label, metric.value])
        summary_sheet.append([])
        summary_sheet.append(["Active alerts", snapshot.active_alerts])
        summary_sheet.append(["Critical alerts", snapshot.critical_alerts])
        summary_sheet.append(["Priority recommendation", snapshot.priority_recommendation])
        summary_sheet.append([])
        summary_sheet.append(["Insight", "Description"])
        for insight in snapshot.insights:
            summary_sheet.append([insight.title, insight.description])

        trend_sheet = workbook.create_sheet("Trend")
        trend_sheet.append(["Bucket", "Revenue", "Charges", "Margin"])
        for bucket, revenue, charges, margin in snapshot.trend_rows:
            trend_sheet.append([bucket, revenue, charges, margin])

        product_sheet = workbook.create_sheet("Top Products")
        product_sheet.append(["Product", "Revenue"])
        for item in snapshot.top_products:
            product_sheet.append([item.name, item.revenue])

        client_sheet = workbook.create_sheet("Top Clients")
        client_sheet.append(["Client", "Revenue"])
        for item in snapshot.top_clients:
            client_sheet.append([item.name, item.revenue])

        alerts_sheet = workbook.create_sheet("Alerts")
        alerts_sheet.append(["Severity", "Rule", "Title", "Recommendation", "Created at"])
        for item in snapshot.alerts_digest:
            alerts_sheet.append(
                [item.severity.upper(), item.ruleName, item.title, item.recommendation, item.createdAt]
            )

        from io import BytesIO

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def build_pdf(self, snapshot: ReportSnapshotData) -> bytes:
        builder = PdfDocumentBuilder()
        margin = 36
        page_width = 595
        page_height = 842
        content_width = page_width - margin * 2
        gutter = 14
        header_height = 116
        header_bottom = page_height - margin - header_height
        metric_card_height = 68
        metric_card_width = (content_width - gutter * 2) / 3
        column_width = (content_width - gutter) / 2
        generated_at = self.format_generated_at(snapshot.generated_at)

        def draw_panel_from_top(x: float, top: float, width: float, height: float) -> None:
            builder.draw_rect(
                x,
                top - height,
                width,
                height,
                fill=PDF_COLORS["white"],
                stroke=PDF_COLORS["border"],
            )

        builder.draw_rect(
            margin,
            header_bottom,
            content_width,
            header_height,
            fill=PDF_COLORS["brand"],
            stroke=PDF_COLORS["brand"],
        )
        builder.draw_rect(
            margin + content_width - 186,
            header_bottom,
            186,
            header_height,
            fill=PDF_COLORS["brand_dark"],
        )

        builder.draw_text(
            "SAL Executive Report",
            margin + 22,
            header_bottom + 82,
            font="F2",
            size=24,
            color=PDF_COLORS["white"],
        )
        builder.draw_text(
            snapshot.company_name,
            margin + 22,
            header_bottom + 54,
            font="F2",
            size=15,
            color=PDF_COLORS["white"],
        )
        builder.draw_text_block(
            snapshot.priority_recommendation,
            margin + 22,
            header_bottom + 36,
            content_width - 230,
            size=10,
            color=PDF_COLORS["white"],
            leading=13,
            max_lines=2,
        )

        meta_top = header_bottom + 86
        for label, value in [
            ("Company", snapshot.company_name),
            ("Reporting period", snapshot.reporting_period),
            ("Generated", generated_at),
        ]:
            builder.draw_text(
                label.upper(),
                margin + content_width - 164,
                meta_top,
                font="F2",
                size=8,
                color=PDF_COLORS["panel"],
            )
            builder.draw_text_block(
                value,
                margin + content_width - 164,
                meta_top - 10,
                142,
                size=10,
                color=PDF_COLORS["white"],
                leading=12,
                max_lines=2,
            )
            meta_top -= 31

        metric_top = header_bottom - 18
        for index, metric in enumerate(snapshot.metrics[:5]):
            row = index // 3
            column = index % 3
            x = margin + column * (metric_card_width + gutter)
            top = metric_top - row * (metric_card_height + 12)

            draw_panel_from_top(x, top, metric_card_width, metric_card_height)
            builder.draw_text(
                metric.label.upper(),
                x + 14,
                top - 21,
                font="F2",
                size=8,
                color=PDF_COLORS["muted"],
            )
            builder.draw_text(
                metric.value,
                x + 14,
                top - 47,
                font="F2",
                size=18,
                color=PDF_COLORS["ink"],
            )

        summary_x = margin + 2 * (metric_card_width + gutter)
        summary_top = metric_top - (metric_card_height + 12)
        draw_panel_from_top(summary_x, summary_top, metric_card_width, metric_card_height)
        builder.draw_text(
            "ACTIVE ALERTS",
            summary_x + 14,
            summary_top - 21,
            font="F2",
            size=8,
            color=PDF_COLORS["muted"],
        )
        builder.draw_text(
            str(snapshot.active_alerts),
            summary_x + 14,
            summary_top - 47,
            font="F2",
            size=18,
            color=PDF_COLORS["ink"],
        )
        builder.draw_right_text(
            f"{snapshot.critical_alerts} critical",
            summary_x + metric_card_width - 14,
            summary_top - 45,
            font="F2",
            size=10,
            color=PDF_COLORS["danger"],
        )

        recommendation_top = metric_top - metric_card_height * 2 - 52
        draw_panel_from_top(margin, recommendation_top, content_width, 74)
        builder.draw_text(
            "PRIORITY FOCUS",
            margin + 16,
            recommendation_top - 20,
            font="F2",
            size=8,
            color=PDF_COLORS["brand_dark"],
        )
        builder.draw_text_block(
            snapshot.priority_recommendation,
            margin + 16,
            recommendation_top - 34,
            content_width - 32,
            font="F2",
            size=13,
            color=PDF_COLORS["ink"],
            leading=16,
            max_lines=3,
        )

        top_sections_top = recommendation_top - 92
        right_column_x = margin + column_width + gutter

        draw_panel_from_top(margin, top_sections_top, column_width, 150)
        builder.draw_text(
            "TOP PRODUCTS",
            margin + 16,
            top_sections_top - 18,
            font="F2",
            size=8,
            color=PDF_COLORS["brand_dark"],
        )
        if snapshot.top_products:
            for index, item in enumerate(snapshot.top_products[:5]):
                row_y = top_sections_top - 46 - index * 20
                builder.draw_text(item.name, margin + 16, row_y, size=10, color=PDF_COLORS["ink"])
                builder.draw_right_text(
                    self.format_currency(item.revenue),
                    margin + column_width - 16,
                    row_y,
                    font="F2",
                    size=10,
                    color=PDF_COLORS["brand_dark"],
                )
        else:
            builder.draw_text_block(
                "No revenue breakdown available.",
                margin + 16,
                top_sections_top - 42,
                column_width - 32,
                size=10,
                color=PDF_COLORS["muted"],
                max_lines=2,
            )

        draw_panel_from_top(right_column_x, top_sections_top, column_width, 150)
        builder.draw_text(
            "TOP CLIENTS",
            right_column_x + 16,
            top_sections_top - 18,
            font="F2",
            size=8,
            color=PDF_COLORS["brand_dark"],
        )
        if snapshot.top_clients:
            for index, item in enumerate(snapshot.top_clients[:5]):
                row_y = top_sections_top - 46 - index * 20
                builder.draw_text(item.name, right_column_x + 16, row_y, size=10, color=PDF_COLORS["ink"])
                builder.draw_right_text(
                    self.format_currency(item.revenue),
                    right_column_x + column_width - 16,
                    row_y,
                    font="F2",
                    size=10,
                    color=PDF_COLORS["brand_dark"],
                )
        else:
            builder.draw_text_block(
                "No revenue breakdown available.",
                right_column_x + 16,
                top_sections_top - 42,
                column_width - 32,
                size=10,
                color=PDF_COLORS["muted"],
                max_lines=2,
            )

        lower_sections_top = top_sections_top - 166
        draw_panel_from_top(margin, lower_sections_top, column_width, 170)
        builder.draw_text(
            "ALERT DIGEST",
            margin + 16,
            lower_sections_top - 18,
            font="F2",
            size=8,
            color=PDF_COLORS["brand_dark"],
        )
        if snapshot.alerts_digest:
            for index, item in enumerate(snapshot.alerts_digest[:4]):
                item_top = lower_sections_top - 34 - index * 32
                builder.draw_rect(
                    margin + 16,
                    item_top - 16,
                    5,
                    18,
                    fill=self.resolve_pdf_severity_color(item.severity),
                )
                builder.draw_text(
                    item.title,
                    margin + 28,
                    item_top - 3,
                    font="F2",
                    size=10,
                    color=PDF_COLORS["ink"],
                )
                builder.draw_text_block(
                    item.recommendation,
                    margin + 28,
                    item_top - 9,
                    column_width - 44,
                    size=8.5,
                    color=PDF_COLORS["muted"],
                    leading=10,
                    max_lines=1,
                )
        else:
            builder.draw_text_block(
                "No active alerts detected.",
                margin + 16,
                lower_sections_top - 42,
                column_width - 32,
                size=10,
                color=PDF_COLORS["muted"],
                max_lines=2,
            )

        draw_panel_from_top(right_column_x, lower_sections_top, column_width, 170)
        builder.draw_text(
            "EXECUTIVE INSIGHTS",
            right_column_x + 16,
            lower_sections_top - 18,
            font="F2",
            size=8,
            color=PDF_COLORS["brand_dark"],
        )
        for index, insight in enumerate(snapshot.insights[:3]):
            item_top = lower_sections_top - 36 - index * 44
            builder.draw_text(
                insight.title,
                right_column_x + 16,
                item_top - 2,
                font="F2",
                size=10,
                color=PDF_COLORS["ink"],
            )
            builder.draw_text_block(
                insight.description,
                right_column_x + 16,
                item_top - 8,
                column_width - 32,
                size=8.5,
                color=PDF_COLORS["muted"],
                leading=10,
                max_lines=2,
            )

        builder.draw_text(
            "Generated by SAL V1",
            margin,
            22,
            size=9,
            color=PDF_COLORS["muted"],
        )
        builder.draw_right_text(generated_at, page_width - margin, 22, size=9, color=PDF_COLORS["muted"])

        return builder.build_bytes()

    def build_output_path(self, export: ReportExportRead) -> Path:
        safe_file_name = export.fileName.replace("/", "-").replace("\\", "-")
        return self.storage_dir / f"{export.id}-{safe_file_name}"

    def format_currency(self, value: float) -> str:
        return f"{value:,.0f} DT"

    def format_percent(self, value: float) -> str:
        return f"{value:.1f}%"

    def format_generated_at(self, value: str) -> str:
        try:
            parsed = datetime.fromisoformat(value)
        except ValueError:
            return value

        return parsed.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    def resolve_pdf_severity_color(self, severity: str) -> PdfColor:
        if severity == "critical":
            return PDF_COLORS["danger"]
        if severity == "warning":
            return PDF_COLORS["warning"]
        return PDF_COLORS["success"]
